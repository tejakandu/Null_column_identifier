"""Excel report and log export."""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink

from . import APP_NAME, AUTHOR, DESIGNED_BY, __version__
from .core import TableResult, format_bytes, format_duration
from .identity import RunIdentity

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF")
NULL_FILL = PatternFill("solid", fgColor="FDE2E1")
WARN_FILL = PatternFill("solid", fgColor="FFF4CE")
OK_FILL = PatternFill("solid", fgColor="E2F0D9")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
LINK_FONT = Font(color="0563C1", underline="single")
TITLE_FONT = Font(bold=True, size=14, color="1F4E79")

OVERVIEW = "Overview"
NULLS = "Null Columns"
RUN_INFO = "Run Info"


@dataclass
class RunInfo:
    executed_by: str = ""
    started: datetime | None = None
    finished: datetime | None = None
    notes: list[str] = field(default_factory=list)
    run_id: str = ""
    identity: RunIdentity | None = None
    query_project: str = ""
    parallel: int = 0
    safety_limit_gb: float = 0.0

    @property
    def who(self) -> str:
        return self.identity.display if self.identity else self.executed_by

    def details(self, results: list[TableResult]) -> list[tuple[str, str]]:
        """Label/value pairs describing the run (Excel, log and PDF share these)."""
        ok = [r for r in results if r.status == "ok"]
        ident = self.identity or RunIdentity()
        rows = [
            ("Tool", f"{APP_NAME} {__version__}"),
            ("Designed by", DESIGNED_BY),
            ("Author", AUTHOR),
            ("Run ID", self.run_id),
            ("Executed by", self.who),
            ("Google account", ident.google_account or "(not available)"),
            ("Sign-in method", ident.auth_method),
            ("Computer", f"{ident.machine} · {ident.os}".strip(" ·")),
            ("Started", self.started.strftime("%Y-%m-%d %H:%M:%S") if self.started else ""),
            ("Finished", self.finished.strftime("%Y-%m-%d %H:%M:%S") if self.finished else ""),
            ("Duration", format_duration((self.finished - self.started).total_seconds())
             if self.started and self.finished else ""),
            ("Query project", self.query_project or "each table's own project"),
            ("Tables at the same time", str(self.parallel) if self.parallel else ""),
            ("Safety limit per query", f"{self.safety_limit_gb:g} GB" if self.safety_limit_gb else "none"),
            ("Tables requested", str(len(results))),
            ("Tables scanned", str(len(ok))),
            ("Tables failed", str(sum(r.status == "error" for r in results))),
            ("Tables cancelled", str(sum(r.status == "cancelled" for r in results))),
            ("Null columns found", str(sum(len(r.null_columns) for r in ok))),
            ("Data scanned", format_bytes(sum(r.bytes_processed for r in results))),
        ]
        return [(k, v) for k, v in rows if v != ""]


def write_report(path: str, results: list[TableResult], info: RunInfo) -> str:
    wb = Workbook()
    overview = wb.active
    overview.title = OVERVIEW
    nulls = wb.create_sheet(NULLS)
    used = {OVERVIEW.lower(), NULLS.lower(), RUN_INFO.lower()}

    detail_names = {}
    for res in results:
        name = _unique_sheet_name(res.target.table, used)
        detail_names[id(res)] = name
        _write_detail(wb.create_sheet(name), res)

    _write_overview(overview, results, detail_names)
    _write_nulls(nulls, results, detail_names)
    _write_run_info(wb.create_sheet(RUN_INFO), results, info)
    wb.save(path)
    return path


# --------------------------------------------------------------------------- sheets


def table_summary(res: TableResult) -> dict:
    """Per-table figures shown in the Overview sheet and stored in run tracking."""
    scanned = res.status == "ok"
    msg = res.error
    if scanned and res.error_columns:
        msg = f"{len(res.error_columns)} column(s) could not be checked"
    if scanned and res.total_rows == 0:
        msg = "Table has no rows, so every column is empty"
    return {
        "project": res.target.project,
        "dataset": res.target.dataset,
        "table": res.target.table,
        "status": _status_text(res),
        "total_rows": res.total_rows,
        "total_columns": len(res.columns) or None,
        "null_columns": len(res.null_columns) if scanned else None,
        "columns_with_data": len(res.data_columns) if scanned else None,
        "coverage": res.coverage,                                   # 0..1
        "not_covered": (1 - res.coverage) if res.coverage is not None else None,
        "message": msg or "",
    }


def _write_overview(ws, results, detail_names):
    headers = ["Project", "Dataset", "Table", "Status", "Total rows", "Total columns",
               "Null columns", "Columns with data", "Coverage", "Not covered",
               "Data scanned", "Details", "Null column names", "Message"]
    ws.append(headers)
    for res in results:
        t = table_summary(res)
        ws.append([
            t["project"], t["dataset"], t["table"], t["status"], t["total_rows"], t["total_columns"],
            t["null_columns"], t["columns_with_data"], t["coverage"], t["not_covered"],
            format_bytes(res.bytes_processed) if res.bytes_processed else "",
            "Open", ", ".join(c.name for c in res.null_columns), t["message"],
        ])
        r = ws.max_row
        _link(ws.cell(r, 12), detail_names[id(res)])
        ws.cell(r, 4).fill = OK_FILL if res.status == "ok" else NULL_FILL
        for col in (9, 10):
            ws.cell(r, col).number_format = "0.00%"
        for col in (5, 6, 7, 8):
            ws.cell(r, col).number_format = "#,##0"
    _finish_table(ws, header_row=1, widths={13: 60, 14: 50})


def _write_nulls(ws, results, detail_names):
    ws.append(["Project", "Dataset", "Table", "Column", "Type", "Mode", "Table rows", "Details"])
    for res in results:
        for col in res.null_columns:
            ws.append([res.target.project, res.target.dataset, res.target.table, col.name,
                       col.spec.data_type, col.spec.mode, res.total_rows, "Open"])
            _link(ws.cell(ws.max_row, 8), detail_names[id(res)])
            ws.cell(ws.max_row, 7).number_format = "#,##0"
    if ws.max_row == 1:
        ws.append(["No null columns were found."])
    _finish_table(ws, header_row=1)


def _write_detail(ws, res: TableResult):
    ws["A1"] = "← Back to Overview"
    _link(ws["A1"], OVERVIEW)
    ws["A2"] = res.target.ref
    ws["A2"].font = TITLE_FONT
    summary = f"Status: {_status_text(res)}"
    if res.total_rows is not None:
        summary += f"   |   Rows: {res.total_rows:,}"
    if res.status == "ok":
        summary += (f"   |   Columns: {len(res.columns)}   |   Null: {len(res.null_columns)}"
                    f"   |   Coverage: {res.coverage:.2%}" if res.coverage is not None else "")
    if res.error:
        summary += f"   |   {res.error}"
    ws["A3"] = summary

    ws.append([])
    ws.append(["Column", "Type", "Mode", "Non-null rows", "Filled", "Result", "Note"])
    header_row = ws.max_row
    rows = res.total_rows or 0
    for col in res.columns:
        if col.non_null is None:
            label, fill = "Not checked", WARN_FILL
        elif col.non_null == 0:
            label, fill = "ALL NULL", NULL_FILL
        else:
            label, fill = "Has data", None
        note = col.error
        if col.spec.mode == "REPEATED" and not note:
            note = "Array: empty arrays count as no value"
        ws.append([col.name, col.spec.data_type, col.spec.mode, col.non_null,
                   (col.non_null / rows) if (rows and col.non_null is not None) else None, label, note])
        r = ws.max_row
        ws.cell(r, 4).number_format = "#,##0"
        ws.cell(r, 5).number_format = "0.00%"
        if fill:
            for c in range(1, 7):
                ws.cell(r, c).fill = fill
    _finish_table(ws, header_row=header_row, widths={1: 45, 7: 45}, measure_from=header_row)


def _write_run_info(ws, results, info: RunInfo):
    rows = info.details(results) + [
        ("Definition", "A column is reported as null when not a single row has a value in it. "
                       "For arrays, an empty array counts as no value. Fields inside repeated "
                       "records are null when no element in any row has a value."),
    ]
    rows += [("Note", n) for n in info.notes]
    ws["A1"] = f"{APP_NAME} — run information"
    ws["A1"].font = TITLE_FONT
    for i, (k, v) in enumerate(rows, start=3):
        ws.cell(i, 1, k).font = Font(bold=True)
        ws.cell(i, 2, v).alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 90


# --------------------------------------------------------------------------- helpers


def _status_text(res: TableResult) -> str:
    return {"ok": "Scanned", "error": "Failed", "cancelled": "Cancelled"}.get(res.status, res.status)


def _link(cell, sheet_name: str):
    quoted = sheet_name.replace("'", "''")
    cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{quoted}'!A1", display=str(cell.value))
    cell.font = LINK_FONT


def _finish_table(ws, header_row: int, widths: dict[int, int] | None = None,
                  measure_from: int = 1):
    for cell in ws[header_row]:
        if cell.value is None:
            continue
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[header_row].height = 22
    for row in ws.iter_rows(min_row=header_row + 1, max_row=ws.max_row):
        for cell in row:
            cell.border = BORDER
    ws.freeze_panes = ws.cell(header_row + 1, 1)
    last_col = get_column_letter(max(1, ws.max_column))
    if ws.max_row > header_row:
        ws.auto_filter.ref = f"A{header_row}:{last_col}{ws.max_row}"

    widths = widths or {}
    for idx in range(1, ws.max_column + 1):
        if idx in widths:
            ws.column_dimensions[get_column_letter(idx)].width = widths[idx]
            continue
        longest = 8
        for (value,) in ws.iter_rows(min_row=measure_from, min_col=idx, max_col=idx, values_only=True):
            if value is not None:
                longest = max(longest, len(str(value)))
        ws.column_dimensions[get_column_letter(idx)].width = min(60, longest + 3)


_INVALID_SHEET_CHARS = re.compile(r"[:\\/?*\[\]]")


def _unique_sheet_name(base: str, used: set[str]) -> str:
    base = _INVALID_SHEET_CHARS.sub("_", base).strip().strip("'") or "Table"
    candidate = base[:31]
    i = 1
    while candidate.lower() in used:
        suffix = f"_{i}"
        candidate = base[:31 - len(suffix)] + suffix
        i += 1
    used.add(candidate.lower())
    return candidate
