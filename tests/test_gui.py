import os

import pytest

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
pytest.importorskip("PySide6")

from openpyxl import load_workbook  # noqa: E402
from PySide6.QtCore import QSettings, Qt  # noqa: E402
from PySide6.QtWidgets import QApplication  # noqa: E402

from nci.core import TableTarget  # noqa: E402
from nci.gui.main_window import MainWindow  # noqa: E402
from nci.gui.style import STYLESHEET  # noqa: E402
from tests.fakes import FakeCatalogClient  # noqa: E402
from tests.test_core import ROWS, SCHEMA  # noqa: E402

TABLES = {
    "proj.sales.orders": (SCHEMA, ROWS),
    "proj.sales.clean": (SCHEMA[:1], ROWS),
    "proj.hr.people": (SCHEMA[:3], ROWS),
}


@pytest.fixture
def app():
    a = QApplication.instance() or QApplication([])
    a.setStyleSheet(STYLESHEET)
    return a


@pytest.fixture
def window(app, tmp_path):
    client = FakeCatalogClient(TABLES)
    settings = QSettings(str(tmp_path / "s.ini"), QSettings.IniFormat)
    w = MainWindow(client_factory=lambda project, key: client, settings=settings)
    w.out_edit.setText(str(tmp_path / "out"))
    w.open_check.setChecked(False)
    w.show()
    yield w
    w.close()


def wait_until(app, cond, timeout=10.0):
    import time
    end = time.monotonic() + timeout
    while time.monotonic() < end:
        app.processEvents()
        if cond():
            return True
        time.sleep(0.01)
    return False


def test_scan_end_to_end(app, window, tmp_path):
    window._add_targets([TableTarget.parse(r) for r in TABLES] + [TableTarget.parse("proj.sales.missing")])
    assert window.model.rowCount() == 4
    window.run_scan()
    assert window._busy()
    assert wait_until(app, lambda: window._last_outcome is not None and window.thread is None)
    out = window._last_outcome
    assert os.path.exists(out.report_path)
    assert os.path.exists(out.log_path)
    states = [r.state for r in window.model.rows]
    assert states == ["nulls", "ok", "nulls", "error"]
    assert window.results_tree.topLevelItemCount() == 4
    assert window.report_btn.isVisible()
    assert not window._busy()
    wb = load_workbook(out.report_path)
    assert "Overview" in wb.sheetnames

    window.grab().save(str(tmp_path / "after_scan.png"))
    os.environ.get("NCI_SCREENSHOT_DIR") and window.grab().save(
        os.path.join(os.environ["NCI_SCREENSHOT_DIR"], "after_scan.png"))


def test_selected_rows_only(app, window):
    window._add_targets([TableTarget.parse(r) for r in TABLES])
    window.view.selectRow(1)
    assert window.run_btn.text() == "Run scan (1)"
    window.run_scan()
    assert wait_until(app, lambda: window._last_outcome is not None and window.thread is None)
    assert [r.state for r in window.model.rows] == ["ready", "ok", "ready"]


def test_cancel(app, window, monkeypatch):
    import nci.core
    window._add_targets([TableTarget.parse(r) for r in TABLES])
    window.run_scan()
    window.cancel_scan()
    assert wait_until(app, lambda: window._last_outcome is not None and window.thread is None)
    assert window._last_outcome.cancelled


def test_settings_and_tables_persist(app, tmp_path):
    settings = QSettings(str(tmp_path / "p.ini"), QSettings.IniFormat)
    w = MainWindow(client_factory=lambda p, k: None, settings=settings)
    w._add_targets([TableTarget.parse("a.b.c")])
    w.limit_spin.setValue(50)
    w.close()
    w2 = MainWindow(client_factory=lambda p, k: None, settings=settings)
    assert w2.model.targets() == [TableTarget("a", "b", "c")]
    assert w2.limit_spin.value() == 50
    w2.close()


def test_edit_cell(app, window):
    window._add_targets([TableTarget.parse("p.d.t")])
    idx = window.model.index(0, 2)
    assert window.model.setData(idx, "t2")
    assert window.model.targets() == [TableTarget("p", "d", "t2")]


def test_parallel_progress_and_filter(app, window):
    targets = [TableTarget("proj", "sales", f"t{i}") for i in range(12)]
    window.pool = type(window.pool)(lambda p, k: FakeCatalogClient(
        {t.ref: (SCHEMA, ROWS) for t in targets}))
    window._add_targets(targets)
    window.list_filter.setText("t1")  # t1, t10, t11
    assert window.proxy.rowCount() == 3
    assert window.run_btn.text() == "Run scan (3)"
    window.parallel_combo.setCurrentIndex(window.parallel_combo.findData(4))
    window.run_scan()
    assert wait_until(app, lambda: window._last_outcome is not None and window.thread is None)
    assert len(window._last_outcome.results) == 3
    assert window.progress.value() == 1000
    assert sum(r.state == "nulls" for r in window.model.rows) == 3
