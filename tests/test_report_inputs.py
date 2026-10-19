from datetime import datetime

from openpyxl import Workbook
from openpyxl import load_workbook

from nci.core import Scanner, TableTarget
from nci.inputs import load_targets, save_targets
from nci.report import RunInfo, write_report
from tests.fakes import FakeClient
from tests.test_core import ROWS, SCHEMA


def _results():
    full = TableTarget("p", "d", "a_very_long_table_name_that_exceeds_excel_limits")
    dup = TableTarget("p", "d2", "a_very_long_table_name_that_exceeds_excel_limits")
    clean = TableTarget("p", "d", "clean")
    client = FakeClient({
        full.ref: (SCHEMA, ROWS),
        dup.ref: (SCHEMA, ROWS),
        clean.ref: (SCHEMA[:1], ROWS),
    })
    scanner = Scanner(lambda t: client)
    missing = TableTarget("p", "d", "missing")
    return [scanner.scan(t) for t in (full, dup, clean, missing)]


def test_report_sheets_and_content(tmp_path):
    path = tmp_path / "r.xlsx"
    write_report(str(path), _results(), RunInfo("tester", datetime(2026, 1, 1), datetime(2026, 1, 1, 0, 1)))
    wb = load_workbook(path)
    assert wb.sheetnames[:2] == ["Overview", "Null Columns"]
    assert wb.sheetnames[-1] == "Run Info"
    detail = wb.sheetnames[2:-1]
    assert len(detail) == 4 and len(set(detail)) == 4
    assert all(len(n) <= 31 for n in detail)

    ov = list(wb["Overview"].iter_rows(values_only=True))
    assert ov[0][:4] == ("Project", "Dataset", "Table", "Status")
    # Fully populated table is still listed (100% coverage).
    clean = next(r for r in ov if r[2] == "clean")
    assert clean[3] == "Scanned" and clean[6] == 0 and clean[8] == 1.0
    missing = next(r for r in ov if r[2] == "missing")
    assert missing[3] == "Failed"
    # Links go to the right sheet.
    link = wb["Overview"].cell(2, 12).hyperlink
    assert link.location == f"'{detail[0]}'!A1"

    nulls = list(wb["Null Columns"].iter_rows(values_only=True))
    assert len(nulls) == 1 + 5 + 5


def test_load_original_layout(tmp_path):
    key = tmp_path / "key.json"
    key.write_text("{}")
    path = tmp_path / "Table_name_credentials.xlsx"
    wb = Workbook()
    wb.active.title = "Tablenames"
    wb.active.append(["x"])
    ws = wb.create_sheet("Credentials")
    ws.append(["CredentialsAccountJson", "Project_id", "Dataset_id", "table_id"])
    ws.append(["key.json", "p1", "d1", "t1"])
    ws.append(["key.json", "p1", "d1", "t2"])
    ws.append([None, "p2", "d2", None])
    wb.save(path)
    targets, warnings = load_targets(str(path))
    assert [t.ref for t in targets] == ["p1.d1.t1", "p1.d1.t2"]
    assert targets[0].key_file == str(key)
    assert warnings == []


def test_load_refs_csv_and_roundtrip(tmp_path):
    csv = tmp_path / "list.csv"
    csv.write_text("p.d.t1\np.d.t2\nbad\n")
    targets, warnings = load_targets(str(csv))
    assert [t.ref for t in targets] == ["p.d.t1", "p.d.t2"]
    assert len(warnings) == 1

    out = tmp_path / "out.xlsx"
    save_targets(str(out), targets)
    again, _ = load_targets(str(out))
    assert again == targets
