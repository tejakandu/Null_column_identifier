"""End-to-end: a scan records who ran it, writes Excel + log + PDF with credits,
tracks the run in BigQuery, and update checks force old versions to close."""

import os
import time

import pytest

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
pytest.importorskip("PySide6")

from openpyxl import load_workbook  # noqa: E402
from PySide6.QtCore import QSettings  # noqa: E402
from PySide6.QtWidgets import QApplication, QLabel  # noqa: E402

from nci import CREDITS, __version__  # noqa: E402
from nci.config import AppConfig  # noqa: E402
from nci.core import TableTarget  # noqa: E402
from nci.gui import main_window as mw  # noqa: E402
from nci.gui import updates  # noqa: E402
from nci.gui.main_window import MainWindow  # noqa: E402
from nci.releases import Release, decide  # noqa: E402
from tests.fakes import FakeCatalogClient  # noqa: E402
from tests.test_core import ROWS, SCHEMA  # noqa: E402

TABLES = {"proj.sales.orders": (SCHEMA, ROWS), "proj.sales.clean": (SCHEMA[:1], ROWS)}
CFG = AppConfig(run_log_table="admin.nci.run_log", releases_table="admin.nci.app_releases",
                download_url="https://intranet/nci.zip")


@pytest.fixture
def app():
    return QApplication.instance() or QApplication([])


def wait_until(app, cond, timeout=15.0):
    end = time.monotonic() + timeout
    while time.monotonic() < end:
        app.processEvents()
        if cond():
            return True
        time.sleep(0.01)
    return False


def make_window(tmp_path, client, cfg=CFG, check_updates=False):
    settings = QSettings(str(tmp_path / "s.ini"), QSettings.IniFormat)
    w = MainWindow(client_factory=lambda p, k: client, settings=settings, config=cfg,
                   check_updates=check_updates)
    w.out_edit.setText(str(tmp_path / "out"))
    w.open_check.setChecked(False)
    w.show()
    return w


def run(app, w):
    w._add_targets([TableTarget.parse(r) for r in TABLES])
    w.run_scan()
    assert wait_until(app, lambda: w._last_outcome is not None and w.thread is None)
    return w._last_outcome


def pdf_text(path):
    """(text, pages), or (None, pages) where QtPdf is not installed.

    QtPdf ships with full PySide6 but not with PySide6-Essentials, which the
    Windows build uses; the page count is then read from the file itself.
    """
    try:
        from PySide6.QtPdf import QPdfDocument
    except ImportError:
        import re

        with open(path, "rb") as fh:
            raw = fh.read()
        assert raw.startswith(b"%PDF")
        return None, len(re.findall(rb"/Type\s*/Page\b", raw))
    doc = QPdfDocument()
    doc.load(path)
    return "\n".join(doc.getAllText(p).text() for p in range(doc.pageCount())), doc.pageCount()


# --------------------------------------------------------------------------- run outputs


def test_scan_writes_credits_identity_pdf_and_tracking(app, tmp_path):
    client = FakeCatalogClient(TABLES, session_user="jane@acme.com")
    w = make_window(tmp_path, client)
    out = run(app, w)

    # Log file: credits and who ran it at the top.
    log = open(out.log_path, encoding="utf-8").read()
    head = log.splitlines()[:4]
    assert CREDITS in head[0] and __version__ in head[0]
    assert head[1] == f"[{head[1][1:9]}] Run ID: {out.run_id}"
    assert "Executed by:" in head[2] and "sign-in: Google login" in head[2]
    assert "Google account: jane@acme.com" in log
    assert "Run recorded for usage tracking" in log

    # PDF log: credits, identity, results, full log, footer on every page.
    assert out.pdf_path.endswith("_log.pdf") and os.path.getsize(out.pdf_path) > 1000
    text, pages = pdf_text(out.pdf_path)
    assert pages >= 1
    if text is not None:
        for needle in ("Designed by", "Teja Kandukuri", "Balasudharam Murugesan", "Executed by",
                       "jane@acme.com", "proj.sales.orders", "items.parts.color", "Full log", out.run_id):
            assert needle in text, needle
        assert text.count("Page ") >= pages and f"Page {pages} of {pages}" in text

    # Excel Run Info sheet.
    info = {r[0]: r[1] for r in load_workbook(out.report_path)["Run Info"].iter_rows(values_only=True) if r[0]}
    assert info["Designed by"] == "Teja Kandukuri" and info["Author"] == "Balasudharam Murugesan"
    assert info["Google account"] == "jane@acme.com" and info["Run ID"] == out.run_id
    assert info["Executed by"]

    # BigQuery run tracking: one row per table, all with this run id.
    assert out.tracked
    (table, rows), = client.inserted
    assert table == "admin.nci.run_log"
    assert sorted(r["table"] for r in rows) == ["clean", "orders"]
    assert all(r["run_id"] == out.run_id and r["google_account"] == "jane@acme.com" for r in rows)
    orders = next(r for r in rows if r["table"] == "orders")
    assert (orders["status"], orders["null_columns"], orders["coverage"]) == ("Scanned", 5, 50.0)
    assert orders["execution_user"] == info["Executed by"]
    assert "Run recorded for usage tracking (2 table row(s))" in log
    assert w.pdf_btn.isVisible() and w.report_btn.isVisible()
    w.close()


def test_tracking_failure_never_breaks_the_run(app, tmp_path):
    client = FakeCatalogClient(TABLES, insert_error=PermissionError("403 Access Denied: Table admin.nci.run_log"),
                               session_user=RuntimeError("no jobs permission"))
    w = make_window(tmp_path, client)
    out = run(app, w)
    assert not out.tracked and out.report_path and out.pdf_path
    log = open(out.log_path, encoding="utf-8").read()
    assert "Run tracking skipped: Access denied" in log
    assert "Google account: not available" in log
    assert "Scanned 2 of 2" in w.status_label.text()
    w.close()


def test_tracking_off_without_config_and_pdf_optional(app, tmp_path):
    client = FakeCatalogClient(TABLES)
    w = make_window(tmp_path, client, cfg=AppConfig())
    w.pdf_check.setChecked(False)
    out = run(app, w)
    assert not out.tracked and client.inserted == []
    assert out.pdf_path == "" and not w.pdf_btn.isVisible()
    assert "Run tracking" not in open(out.log_path, encoding="utf-8").read()
    w.close()


# --------------------------------------------------------------------------- credits in the UI


def test_header_shows_credits_and_about(app, tmp_path, monkeypatch):
    w = make_window(tmp_path, FakeCatalogClient(TABLES))
    labels = " ".join(lbl.text() for lbl in w.findChildren(QLabel))
    assert "Teja Kandukuri" in labels and "Balasudharam Murugesan" in labels
    shown = []
    monkeypatch.setattr(updates.AboutDialog, "exec",
                        lambda self: shown.append(" ".join(lbl.text() for lbl in self.findChildren(QLabel))))
    w.show_about()
    assert "Teja Kandukuri" in shown[0] and __version__ in shown[0]
    w.close()


# --------------------------------------------------------------------------- updates


def newer(mandatory, version="9.0.0", url="https://dl/9.0.0", notes="Big update"):
    return {"version": version, "mandatory": mandatory, "download_url": url, "release_notes": notes}


def test_forced_update_closes_the_app(app, tmp_path, monkeypatch):
    client = FakeCatalogClient(TABLES, releases=[newer(True)])
    shown = []
    monkeypatch.setattr(mw.ForcedUpdateDialog, "exec", lambda self: shown.append(self.url))
    w = make_window(tmp_path, client)
    w.updates.check()
    assert wait_until(app, lambda: shown)
    assert shown == ["https://dl/9.0.0"]
    assert w._force_quit and not w.isVisible()


def test_forced_update_waits_for_running_scan(app, tmp_path, monkeypatch):
    client = FakeCatalogClient(TABLES)
    shown = []
    w = make_window(tmp_path, client)
    # Record the app's state at the moment the prompt opens.
    monkeypatch.setattr(mw.ForcedUpdateDialog, "exec",
                        lambda self: shown.append((w._busy(), w._last_outcome is not None)))
    w._add_targets([TableTarget.parse(r) for r in TABLES])
    w.run_scan()
    assert w._busy()
    w._on_update_decision(decide([Release("9.0.0", mandatory=True)]))  # arrives mid-scan
    assert shown == []                      # not while the scan runs
    assert "will close when this scan ends" in w.log_view.toPlainText()
    assert wait_until(app, lambda: shown)
    assert shown == [(False, True)]         # only after the scan finished and saved
    assert os.path.exists(w._last_outcome.report_path)
    assert not w.isVisible()


def test_optional_update_only_shows_button(app, tmp_path, monkeypatch):
    client = FakeCatalogClient(TABLES, releases=[newer(False, "9.1.0")])
    monkeypatch.setattr(mw.ForcedUpdateDialog, "exec", lambda self: pytest.fail("must not force"))
    w = make_window(tmp_path, client)
    w.updates.check()
    assert wait_until(app, lambda: w.update_btn.isVisible())
    assert w.update_btn.text() == "Update to 9.1.0"
    out = run(app, w)  # the app stays usable
    assert out.report_path
    w.close()


def test_update_check_failure_fails_open(app, tmp_path, monkeypatch):
    class Broken(FakeCatalogClient):
        def list_rows(self, *a, **k):
            raise PermissionError("403 Access Denied: releases")

    monkeypatch.setattr(mw.ForcedUpdateDialog, "exec", lambda self: pytest.fail("must not force"))
    w = make_window(tmp_path, Broken(TABLES))
    w.updates.check()
    assert wait_until(app, lambda: "Update check skipped" in w.log_view.toPlainText())
    assert w._forced_decision is None and not w.update_btn.isVisible()
    assert run(app, w).report_path
    w.close()


def test_no_checks_without_config(app, tmp_path):
    client = FakeCatalogClient(TABLES, releases=[newer(True)])
    w = make_window(tmp_path, client, cfg=AppConfig(), check_updates=True)
    assert not w.updates.enabled
    w.updates.check()
    app.processEvents()
    assert w._forced_decision is None
    w.close()


def test_forced_dialog_cannot_be_dismissed(app, monkeypatch):
    opened = []
    monkeypatch.setattr(updates.QDesktopServices, "openUrl", lambda url: opened.append(url.toString()))
    decision = decide([Release("9.0.0", mandatory=True, download_url="https://dl/9", release_notes="n")])
    dlg = updates.ForcedUpdateDialog(decision, CFG)
    assert dlg.download_btn.isEnabled()
    dlg.reject()  # Esc / Alt+F4 behaves like Quit: the app still closes
    assert dlg.result() == dlg.DialogCode.Accepted and not dlg.downloaded
    dlg2 = updates.ForcedUpdateDialog(decision, CFG)
    dlg2._download()
    assert opened == ["https://dl/9"] and dlg2.downloaded

    # Falls back to the configured link, or disables Download if there is none.
    no_url = decide([Release("9.0.0", mandatory=True)])
    assert updates.ForcedUpdateDialog(no_url, CFG).url == "https://intranet/nci.zip"
    assert not updates.ForcedUpdateDialog(no_url, AppConfig()).download_btn.isEnabled()
