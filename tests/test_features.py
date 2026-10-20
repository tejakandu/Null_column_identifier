"""Credits, who-ran-it, run tracking, releases/forced updates, admin tool, config."""

import json
from datetime import datetime, timezone
from types import SimpleNamespace

import pytest

from nci import AUTHOR, CREDITS, DESIGNED_BY, __version__
from nci import admin, config
from nci.core import Scanner, TableTarget
from nci.identity import RunIdentity, google_account, local_identity
from nci.releases import Release, decide, fetch_releases, parse_version
from nci.tracking import RUN_LOG_SCHEMA, build_rows, record_run
from tests.fakes import FakeClient
from tests.test_core import ROWS, SCHEMA


def _results():
    good = TableTarget("p", "d", "orders")
    clean = TableTarget("p", "d", "clean")
    client = FakeClient({good.ref: (SCHEMA, ROWS), clean.ref: (SCHEMA[:1], ROWS)})
    scanner = Scanner(lambda t: client)
    return [scanner.scan(good), scanner.scan(clean), scanner.scan(TableTarget("p", "d", "missing"))]


IDENT = RunIdentity("Jane Doe", "jdoe", "LAPTOP-42", "Windows 11", "jane@acme.com", "Google login")


# --------------------------------------------------------------------------- credits


def test_credits_text():
    assert DESIGNED_BY == "Teja Kandukuri"
    assert AUTHOR == "Balasudharam Murugesan"
    assert CREDITS == "Designed by Teja Kandukuri · Author: Balasudharam Murugesan"


# --------------------------------------------------------------------------- identity


def test_identity_display():
    assert IDENT.display == "Jane Doe (jdoe) on LAPTOP-42"
    assert RunIdentity(login="jdoe", machine="PC").display == "jdoe on PC"
    assert RunIdentity(full_name="jdoe", login="jdoe").display == "jdoe"
    assert RunIdentity().display == "unknown user"


def test_local_identity_and_session_user():
    ident = local_identity("Google login")
    assert ident.login and ident.machine and ident.os and ident.auth_method == "Google login"
    client = FakeClient({}, session_user="sa@proj.iam.gserviceaccount.com")
    assert google_account(client) == "sa@proj.iam.gserviceaccount.com"
    assert "SESSION_USER()" in client.queries[0]


# --------------------------------------------------------------------------- tracking

COLUMNS = ["run_id", "execution_time", "execution_user", "project", "dataset", "table", "status",
           "total_rows", "total_columns", "null_columns", "columns_with_data", "coverage",
           "not_covered", "message", "google_account", "app_version"]


def test_schema_has_the_requested_columns_in_order():
    assert [c[0] for c in RUN_LOG_SCHEMA] == COLUMNS


def test_rows_one_per_table_match_schema():
    started = datetime(2026, 10, 20, 9, 0, 0)
    rows = build_rows(run_id="r1", results=_results(), identity=IDENT, started=started)
    assert len(rows) == 3
    json.dumps(rows)  # streaming insert sends JSON
    by_table = {r["table"]: r for r in rows}
    for r in rows:
        assert list(r) == COLUMNS
        assert r["run_id"] == "r1" and r["app_version"] == __version__
        assert r["execution_user"] == "Jane Doe (jdoe) on LAPTOP-42"
        assert r["google_account"] == "jane@acme.com"
        assert datetime.fromisoformat(r["execution_time"]).tzinfo is not None  # UTC, not naive
        for name, _t, mode, _d in RUN_LOG_SCHEMA:
            if mode == "REQUIRED":
                assert r[name] is not None, name

    o = by_table["orders"]
    assert (o["project"], o["dataset"], o["status"]) == ("p", "d", "Scanned")
    assert (o["total_rows"], o["total_columns"], o["null_columns"], o["columns_with_data"]) == (2, 10, 5, 5)
    assert (o["coverage"], o["not_covered"]) == (50.0, 50.0)
    assert o["message"] is None
    c = by_table["clean"]
    assert (c["coverage"], c["not_covered"], c["null_columns"]) == (100.0, 0.0, 0)
    m = by_table["missing"]
    assert m["status"] == "Failed" and m["message"].startswith("Not found")
    assert m["total_rows"] is None and m["coverage"] is None


def test_rows_agree_with_excel_overview(tmp_path):
    from openpyxl import load_workbook

    from nci.report import RunInfo, write_report

    results = _results()
    write_report(str(tmp_path / "r.xlsx"), results, RunInfo())
    sheet = list(load_workbook(tmp_path / "r.xlsx")["Overview"].iter_rows(values_only=True))
    head, body = sheet[0], sheet[1:]
    rows = build_rows(run_id="r", results=results, identity=IDENT, started=datetime.now())
    for excel, row in zip(body, rows):
        e = dict(zip(head, excel))
        assert (e["Project"], e["Dataset"], e["Table"], e["Status"]) == \
            (row["project"], row["dataset"], row["table"], row["status"])
        assert e["Total rows"] == row["total_rows"] and e["Null columns"] == row["null_columns"]
        assert e["Columns with data"] == row["columns_with_data"]
        if e["Coverage"] is not None:
            assert round(e["Coverage"] * 100, 2) == row["coverage"]
        assert (e["Message"] or None) == row["message"]


def test_cancelled_and_unknown_user():
    from nci.core import TableResult

    rows = build_rows(run_id="r", identity=RunIdentity(), started=datetime.now(),
                      results=[TableResult(TableTarget("p", "d", "t"), status="cancelled",
                                           error="Cancelled before it started")])
    assert rows[0]["status"] == "Cancelled" and rows[0]["execution_user"] is None
    assert rows[0]["message"] == "Cancelled before it started"


def test_record_run_batches_and_dedupes():
    client = FakeClient({})
    rows = [{"run_id": "abc", "n": i} for i in range(1203)]
    record_run(client, "a.b.run_log", rows)
    assert [len(b) for _t, b in client.inserted] == [500, 500, 203]
    assert sum((b for _t, b in client.inserted), []) == rows
    with pytest.raises(RuntimeError, match="rejected"):
        record_run(FakeClient({}, insert_error="no such field"), "a.b.c", [{"run_id": "x"}])


# --------------------------------------------------------------------------- releases


@pytest.mark.parametrize("text,expected", [
    ("2.2.0", (2, 2)), ("v2.10.1", (2, 10, 1)), ("2.3.0-beta", (2, 3)), ("3", (3,)), ("", (0,)),
])
def test_parse_version(text, expected):
    assert parse_version(text) == expected


def test_version_ordering():
    assert parse_version("2.10.0") > parse_version("2.9.9")
    assert parse_version("2.2") == parse_version("2.2.0")


def R(v, mandatory=False, notes=""):
    return Release(v, mandatory=mandatory, download_url=f"https://dl/{v}", release_notes=notes)


def test_decide():
    assert decide([], "2.2.0").kind == "none"
    assert decide([R("2.1.0", True), R("2.2.0")], "2.2.0").kind == "none"  # old mandatory: ignored
    d = decide([R("2.2.1", notes="fix"), R("2.3.0", notes="new")], "2.2.0")
    assert d.kind == "optional" and d.latest.version == "2.3.0"
    assert d.notes == ("2.3.0: new", "2.2.1: fix")
    d = decide([R("2.2.1", True), R("2.3.0")], "2.2.0")  # a mandatory step in between
    assert d.kind == "forced" and d.required.version == "2.2.1" and d.latest.version == "2.3.0"
    assert decide([R("2.2.1", True)], "2.2.1").kind == "none"  # already on it


def test_fetch_releases_skips_withdrawn_and_blank():
    when = datetime(2026, 10, 1, tzinfo=timezone.utc)
    client = FakeClient({}, releases=[
        {"version": "2.3.0", "released_at": when, "mandatory": True, "download_url": "u", "withdrawn": True},
        {"version": "2.2.5", "released_at": when, "mandatory": None, "download_url": None},
        {"version": "", "mandatory": True},
    ])
    rel = fetch_releases(client, "a.b.c")
    assert [r.version for r in rel] == ["2.2.5"]
    assert rel[0].mandatory is False and rel[0].download_url == ""


# --------------------------------------------------------------------------- config


def test_config_loading(tmp_path, monkeypatch):
    path = tmp_path / "c.json"
    path.write_text(json.dumps({"run_log_table": "a.b.runs", "releases_table": "a.b.rel",
                                "download_url": "https://x", "update_check_minutes": 1}))
    monkeypatch.setenv("NCI_CONFIG", str(path))
    cfg = config.load_config()
    assert cfg.tracking_enabled and cfg.updates_enabled and cfg.source == str(path)
    assert cfg.update_check_minutes == 5  # clamped: never hammer BigQuery
    path.write_text("{broken")
    assert config.load_config().source != str(path)  # falls back to the bundled file


def test_override_cannot_blank_bundled_settings(tmp_path, monkeypatch):
    bundled = tmp_path / "pkg"
    bundled.mkdir()
    (bundled / "nci_config.json").write_text(json.dumps(
        {"run_log_table": "a.b.runs", "releases_table": "a.b.rel", "download_url": "https://x"}))
    override = tmp_path / "override.json"
    override.write_text(json.dumps({"releases_table": "", "run_log_table": None,
                                    "download_url": "https://new"}))
    monkeypatch.setattr(config, "__file__", str(bundled / "config.py"))
    monkeypatch.setenv("NCI_CONFIG", str(override))
    cfg = config.load_config()
    assert cfg.updates_enabled and cfg.tracking_enabled      # blanks ignored
    assert cfg.download_url == "https://new"                  # real values still override


def test_bundled_config_is_valid_and_off_by_default(monkeypatch):
    monkeypatch.delenv("NCI_CONFIG", raising=False)
    cfg = config.load_config()
    assert cfg.source.endswith("nci_config.json")
    assert not cfg.tracking_enabled and not cfg.updates_enabled


# --------------------------------------------------------------------------- admin tool


def test_ddl_is_valid_bigquery_sql():
    sqlglot = pytest.importorskip("sqlglot")
    sql = admin.ddl("admin-proj.nci")
    statements = [s for s in sqlglot.parse(sql, read="bigquery") if s is not None]
    assert len(statements) == 2
    assert "PARTITION BY DATE(execution_time)" in sql
    assert "CLUSTER BY project, dataset, `table`" in sql


class AdminClient(FakeClient):
    def __init__(self, releases=()):
        super().__init__({}, releases=releases)
        self.jobs = []

    def query(self, sql, job_config=None):
        self.jobs.append((sql, {p.name: p.value for p in job_config.query_parameters}))
        return SimpleNamespace(result=lambda: None, num_dml_affected_rows=1)


def test_admin_release(monkeypatch, capsys):
    client = AdminClient(releases=[{"version": "2.2.0", "mandatory": False}])
    monkeypatch.setattr(admin, "_client", lambda project, key: client)
    rc = admin.main(["release", "2.3.0", "--url", "https://dl/2.3.0", "--notes", "Faster", "--mandatory",
                     "--table", "a.b.app_releases", "--by", "Teja"])
    assert rc == 0
    sql, params = client.jobs[0]
    assert sql.startswith("INSERT INTO `a.b.app_releases`")
    assert params == {"v": "2.3.0", "m": True, "u": "https://dl/2.3.0", "n": "Faster", "by": "Teja"}
    assert "MANDATORY" in capsys.readouterr().out

    # Refuses to publish a version that is not newer, unless forced.
    assert admin.main(["release", "2.1.0", "--url", "u", "--table", "a.b.c"]) == 1
    assert admin.main(["release", "2.1.0", "--url", "u", "--table", "a.b.c", "--force"]) == 0


def test_admin_withdraw_and_list(monkeypatch, capsys):
    client = AdminClient(releases=[{"version": "2.3.0", "mandatory": True, "release_notes": "n"}])
    monkeypatch.setattr(admin, "_client", lambda project, key: client)
    assert admin.main(["withdraw", "2.3.0", "--table", "a.b.c"]) == 0
    assert client.jobs[0][0].startswith("UPDATE `a.b.c` SET withdrawn = TRUE")
    assert admin.main(["list", "--table", "a.b.c"]) == 0
    out = capsys.readouterr().out
    assert "2.3.0" in out and "MANDATORY" in out


def test_admin_setup_writes_config(monkeypatch, tmp_path):
    from google.cloud import bigquery

    from google.api_core.exceptions import NotFound

    created = []

    def missing(ref):
        raise NotFound(ref)

    client = SimpleNamespace(get_dataset=missing, create_dataset=lambda ds: created.append(ds.dataset_id),
                             create_table=lambda t: created.append(t.table_id))
    monkeypatch.setattr(admin, "_client", lambda project, key: client)
    cfg = tmp_path / "nci_config.json"
    cfg.write_text(json.dumps({"_help": "keep me", "download_url": "https://dl"}))
    monkeypatch.setattr(admin, "CONFIG_PATH", str(cfg))
    assert admin.main(["setup", "--dataset", "adm.nci"]) == 0
    assert created == ["nci", "run_log", "app_releases"]
    data = json.loads(cfg.read_text())
    assert data == {"_help": "keep me", "download_url": "https://dl",
                    "run_log_table": "adm.nci.run_log", "releases_table": "adm.nci.app_releases"}
    assert isinstance(bigquery.Table("a.b.c"), bigquery.Table)


def test_admin_setup_uses_existing_dataset_without_creating(monkeypatch, tmp_path, capsys):
    made = []

    def no_create(ds):
        raise AssertionError("must not try to create an existing dataset")

    client = SimpleNamespace(get_dataset=lambda ref: SimpleNamespace(location="US"),
                             create_dataset=no_create, create_table=lambda t: made.append(t.table_id))
    monkeypatch.setattr(admin, "_client", lambda project, key: client)
    monkeypatch.setattr(admin, "CONFIG_PATH", str(tmp_path / "c.json"))
    assert admin.main(["setup", "--dataset", "dp-qa-derived-e0f9.Qa_automation"]) == 0
    assert made == ["run_log", "app_releases"]
    out = capsys.readouterr().out
    assert "Using existing dataset dp-qa-derived-e0f9.Qa_automation" in out
    cfg = json.loads((tmp_path / "c.json").read_text())
    assert cfg["run_log_table"] == "dp-qa-derived-e0f9.Qa_automation.run_log"
    assert cfg["releases_table"] == "dp-qa-derived-e0f9.Qa_automation.app_releases"
